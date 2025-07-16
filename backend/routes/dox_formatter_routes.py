from fastapi import APIRouter
from schemas.docx_model import DocumentRequest
from fastapi.responses import StreamingResponse
import io
from docx import Document as DocxDocument

router = APIRouter(prefix="/api/v1/docx", tags=["docx"])

@router.post("/format-and-download")
async def create_document_from_text(request: DocumentRequest):
    """
    Accepts raw text in a JSON request, formats it into a .docx file in memory,
    and returns it for download.
    """
    doc = DocxDocument()
    
    # --- Formatting Logic ---
    # This logic parses the raw_text and applies styles.
    # It assumes markdown-like headings (#, ##) and lists (-).
    for line in request.raw_text.strip().split('\n'):
        stripped_line = line.strip()
        if not stripped_line:
            continue

        if stripped_line.startswith('#### '):
            doc.add_heading(stripped_line.replace('#### ', ''), level=4)
        elif stripped_line.startswith('### '):
            doc.add_heading(stripped_line.replace('### ', ''), level=3)
        elif stripped_line.startswith('## '):
            doc.add_heading(stripped_line.replace('## ', ''), level=2)
        elif stripped_line.startswith('# '):
            doc.add_heading(stripped_line.replace('# ', ''), level=1)
        elif stripped_line.startswith('- '):
            # Handles list items and bold text within them
            p = doc.add_paragraph(style='List Bullet')
            clean_line = stripped_line[2:] # Remove '- '
            parts = clean_line.split('**')
            for i, part in enumerate(parts):
                if i % 2 == 1:
                    p.add_run(part).bold = True
                else:
                    p.add_run(part)
        else:
            # Handles regular paragraphs with bold text
            p = doc.add_paragraph()
            parts = stripped_line.split('**')
            for i, part in enumerate(parts):
                if i % 2 == 1:
                    p.add_run(part).bold = True
                else:
                    p.add_run(part)

    # --- In-Memory File Handling ---
    # 1. Create an in-memory binary stream.
    file_stream = io.BytesIO()

    # 2. Save the document to the in-memory stream.
    doc.save(file_stream)

    # 3. Important: Go back to the beginning of the stream.
    file_stream.seek(0)
    
    # --- Custom Headers ---
    # The Content-Disposition header is what tells the browser to download the file
    # and what to name it.
    headers = {
        'Content-Disposition': f'attachment; filename="{request.file_name}"'
    }

    # --- Return StreamingResponse ---
    # Stream the content of the in-memory file.
    return StreamingResponse(
        content=file_stream,
        media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        headers=headers
    )
