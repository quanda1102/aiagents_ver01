import os
from sqlalchemy.orm import Session
from fastapi import HTTPException
from openai import OpenAI
from docx import Document
import PyPDF2
from config import config
from schemas.lecture import LectureInput, LectureContext, LectureStructure, LectureOutput
from models.lecture import Lecture

client = OpenAI(api_key=config.OPENAI_API_KEY)

class LectureService:
    @staticmethod
    def extract_file_content(file_path: str) -> str:
        try:
            if file_path.endswith(".docx"):
                doc = Document(file_path)
                return "\n".join([para.text for para in doc.paragraphs])
            elif file_path.endswith(".pdf"):
                with open(file_path, "rb") as file:
                    reader = PyPDF2.PdfReader(file)
                    return "\n".join([page.extract_text() for page in reader.pages])
            elif file_path.endswith(".txt"):
                with open(file_path, "r", encoding="utf-8") as file:
                    return file.read()
            raise HTTPException(status_code=400, detail="Unsupported file format")
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Error extracting file: {str(e)}")

    @staticmethod
    async def collect_context(input: LectureInput, db: Session) -> LectureContext:
        try:
            raw_content = input.raw_content
            if os.path.exists(raw_content):
                raw_content = LectureService.extract_file_content(raw_content)

            prompt = f"""
Phân tích nội dung sau và trích xuất các thành phần:
- Mục tiêu bài học
- Kỹ năng cần đạt
- Hoạt động chính
- Thiết bị dạy học
- Đánh giá
Nội dung: {raw_content}
Khối lớp: {input.grade_level}
Môn học: {input.subject}
Trả về định dạng JSON:
{{
    "objectives": [],
    "skills": [],
    "activities": [],
    "equipment": [],
    "assessment": []
}}
"""
            response = client.chat.completions.create(
                model="gpt-4o-mini",
                messages=[{"role": "user", "content": prompt}],
                max_tokens=1000,
                temperature=0.5
            )

            context = response.choices[0].message.content.strip()

            if context.startswith("```json"):
                context = context[len("```json"):].strip()
            if context.startswith("```"):
                context = context[len("```"):].strip()
            if context.endswith("```"):
                context = context[:-3].strip()

            return LectureContext.parse_raw(context)
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Error in collect_context: {str(e)}")

    @staticmethod
    async def structure_lecture(context: LectureContext, input: LectureInput) -> LectureStructure:
        try:
            prompt = f"""
Dựa trên context sau, tạo khung bài giảng theo chuẩn Bộ GD Việt Nam:
- Mục tiêu bài học
- Chuẩn kiến thức
- Hoạt động: khởi động, hình thành kiến thức, luyện tập, vận dụng
- Câu hỏi
- Thiết bị dạy học
- Đánh giá
Context: {context.dict()}
Khối lớp: {input.grade_level}
Môn học: {input.subject}
Trả về định dạng JSON:
{{
    "title": "",
    "objectives": [],
    "knowledge_standards": [],
    "activities": {{
        "warm_up": "",
        "knowledge_formation": "",
        "practice": "",
        "application": ""
    }},
    "questions": [],
    "equipment": [],
    "assessment": []
}}
"""
            response = client.chat.completions.create(
                model="gpt-4o-mini",
                messages=[{"role": "user", "content": prompt}],
                max_tokens=1500,
                temperature=0.5
            )

            structure = response.choices[0].message.content.strip()

            if structure.startswith("```json"):
                structure = structure[len("```json"):].strip()
            if structure.startswith("```"):
                structure = structure[len("```"):].strip()
            if structure.endswith("```"):
                structure = structure[:-3].strip()

            return LectureStructure.parse_raw(structure)
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Error in structure_lecture: {str(e)}")

    @staticmethod
    async def write_content(structure: LectureStructure, section: str = None) -> LectureStructure:
        try:
            if section:
                prompt = f"""
Viết nội dung chi tiết cho phần {section} của bài giảng:
Structure: {structure.dict()}
Trả về định dạng JSON với nội dung chi tiết cho phần được yêu cầu.
"""
            else:
                prompt = f"""
Viết nội dung chi tiết cho toàn bộ bài giảng:
Structure: {structure.dict()}
Trả về định dạng JSON với nội dung đầy đủ.
"""

            response = client.chat.completions.create(
                model="gpt-4o-mini",
                messages=[{"role": "user", "content": prompt}],
                max_tokens=2000,
                temperature=0.7
            )

            content = response.choices[0].message.content.strip()

            if content.startswith("```json"):
                content = content[len("```json"):].strip()
            if content.startswith("```"):
                content = content[len("```"):].strip()
            if content.endswith("```"):
                content = content[:-3].strip()

            return LectureStructure.parse_raw(content)
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Error in write_content: {str(e)}")

    @staticmethod
    async def edit_lecture(structure: LectureStructure, edit_request: str) -> LectureStructure:
        try:
            prompt = f"""
Chỉnh sửa bài giảng theo yêu cầu: {edit_request}
Structure: {structure.dict()}
Trả về định dạng JSON với bài giảng đã chỉnh sửa.
"""
            response = client.chat.completions.create(
                model="gpt-4o-mini",
                messages=[{"role": "user", "content": prompt}],
                max_tokens=1500,
                temperature=0.7
            )

            edited = response.choices[0].message.content.strip()

            if edited.startswith("```json"):
                edited = edited[len("```json"):].strip()
            if edited.startswith("```"):
                edited = edited[len("```"):].strip()
            if edited.endswith("```"):
                edited = edited[:-3].strip()

            return LectureStructure.parse_raw(edited)
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Error in edit_lecture: {str(e)}")

    @staticmethod
    def save_lecture(structure: LectureStructure, input: LectureInput, teacher_id: int, db: Session) -> LectureOutput:
        try:
            lecture = Lecture(
                teacher_id=teacher_id,
                title=input.title,
                content=structure.dict(),
                grade_level=input.grade_level,
                subject=input.subject
            )
            db.add(lecture)
            db.commit()
            db.refresh(lecture)
            return LectureOutput.from_orm(lecture)
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Error saving lecture: {str(e)}")

    @staticmethod
    def export_lecture_to_docx(lecture: LectureOutput) -> str:
        try:
            doc = Document()
            doc.add_heading(lecture.title, 0)
            for section, content in lecture.content["activities"].items():
                doc.add_heading(section.replace("_", " ").title(), level=1)
                doc.add_paragraph(content)
            file_path = f"lectures/{lecture.id}_{lecture.title}.docx"
            doc.save(file_path)
            return file_path
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Error exporting lecture: {str(e)}")
