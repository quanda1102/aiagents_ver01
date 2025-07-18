import pytest
from fastapi.testclient import TestClient
from main import app
import io
from docx import Document as DocxDocument

client = TestClient(app)

# --- Test Cases for DOCX Formatter ---

def test_format_document_with_all_elements():
    """
    Tests the successful creation of a .docx file with headings, paragraphs, lists, and tables.
    """
    raw_text_payload = """
# Main Title

This is the first paragraph. It contains **bold text**.

## Subheading

- List item 1
- List item 2 with **bolding**

[TABLE_START]
Header 1 | Header 2
---|---
Row 1, Col 1 | Row 1, Col 2
Row 2, Col 1 | Row 2, Col 2
[TABLE_END]

Another paragraph after the table.
"""
    
    request_data = {
        "raw_text": raw_text_payload,
        "file_name": "comprehensive_document.docx"
    }
    
    response = client.post("/api/v1/docx/format-and-download", json=request_data)
    
    # 1. Check HTTP response
    assert response.status_code == 200
    assert response.headers['content-type'] == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    assert "comprehensive_document.docx" in response.headers['content-disposition']
    
    # 2. Verify the DOCX content itself
    document_stream = io.BytesIO(response.content)
    doc = DocxDocument(document_stream)
    
    # Check headings
    assert doc.paragraphs[0].text == "Main Title"
    assert doc.paragraphs[0].style.name == 'Heading 1'
    assert doc.paragraphs[2].text == "Subheading"
    assert doc.paragraphs[2].style.name == 'Heading 2'
    
    # Check list
    assert doc.paragraphs[4].text == "List item 1"
    assert doc.paragraphs[4].style.name == 'List Bullet'
    assert doc.paragraphs[5].text == "List item 2 with bolding"
    
    # Check table
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3
    assert len(table.columns) == 2
    assert table.cell(0, 0).text == "Header 1"
    assert table.cell(1, 1).text == "Row 1, Col 2"
    
    # Check final paragraph
    # The index might change depending on how paragraphs are added around the table
    assert doc.paragraphs[-1].text == "Another paragraph after the table."

def test_format_document_validation_error():
    """
    Tests that the endpoint returns a 422 Unprocessable Entity error
    if the required 'raw_text' field is missing.
    """
    request_data = {
        "file_name": "invalid_request.docx"
        # Missing "raw_text"
    }
    
    response = client.post("/api/v1/docx/format-and-download", json=request_data)
    
    assert response.status_code == 422

def test_format_document_with_empty_text():
    """
    Tests that providing an empty string for 'raw_text' results in a valid,
    but empty, .docx file.
    """
    request_data = {
        "raw_text": "",
        "file_name": "empty_document.docx"
    }
    
    response = client.post("/api/v1/docx/format-and-download", json=request_data)
    
    assert response.status_code == 200
    assert "empty_document.docx" in response.headers['content-disposition']
    
    # Verify it's a valid docx file with no content (or just a single empty paragraph)
    document_stream = io.BytesIO(response.content)
    doc = DocxDocument(document_stream)
    # A new document usually has one empty paragraph
    assert len(doc.paragraphs) <= 1
    if len(doc.paragraphs) == 1:
        assert doc.paragraphs[0].text == ""
    assert len(doc.tables) == 0
